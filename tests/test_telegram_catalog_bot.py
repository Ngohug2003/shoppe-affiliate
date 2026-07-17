import json
from io import BytesIO
from unittest.mock import AsyncMock

import httpx
import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]

from app.schemas.responses.affiliate_catalog import AffiliateProductResponse
from app.schemas.telegram import (
    TelegramChat,
    TelegramDocument,
    TelegramMessage,
    TelegramUpdate,
)
from app.services.telegram_catalog_bot import (
    CatalogApiClient,
    TelegramCatalogBot,
)


async def test_catalog_client_authenticates_and_imports_product() -> None:
    requested_paths: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested_paths.append(request.url.path)
        if request.url.path == "/api/v1/auth/token":
            return httpx.Response(
                200,
                json={
                    "status": {"code": 200, "message": "Đăng nhập thành công"},
                    "data": {"access_token": "token", "token_type": "bearer"},
                },
            )
        assert request.headers["Authorization"] == "Bearer token"
        assert json.loads(request.content) == {"url": "https://vn.shp.ee/test"}
        return httpx.Response(
            201,
            json={
                "status": {
                    "code": 201,
                    "message": "Thêm sản phẩm affiliate thành công",
                },
                "data": {
                    "shop_id": "123",
                    "item_id": "456",
                    "title": "Sản phẩm",
                    "image_url": "https://down-vn.img.susercontent.com/file/image",
                    "product_url": "https://shopee.vn/product/123/456",
                    "affiliate_url": "https://s.shopee.vn/an_redir?affiliate_id=1",
                    "metadata_source": "shopee_open_graph",
                },
            },
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        catalog = CatalogApiClient(
            base_url="http://api:8000",
            admin_email="admin@example.com",
            admin_password="password",
            client=client,
        )
        result = await catalog.import_product("https://vn.shp.ee/test")

    assert result.item_id == "456"
    assert requested_paths == ["/api/v1/auth/token", "/api/v1/affiliate-products"]


async def test_telegram_message_calls_catalog_and_replies() -> None:
    sent_payloads: list[dict[str, object]] = []

    def telegram_handler(request: httpx.Request) -> httpx.Response:
        sent_payloads.append(json.loads(request.content))
        return httpx.Response(200, json={"ok": True})

    catalog = AsyncMock(spec=CatalogApiClient)
    catalog.import_product.return_value = AffiliateProductResponse(
        shop_id="123",
        item_id="456",
        title="Sản phẩm",
        image_url="https://down-vn.img.susercontent.com/file/image",
        product_url="https://shopee.vn/product/123/456",
        affiliate_url="https://s.shopee.vn/an_redir?affiliate_id=1",
        metadata_source="shopee_open_graph",
    )
    async with httpx.AsyncClient(
        transport=httpx.MockTransport(telegram_handler)
    ) as telegram_client:
        bot = TelegramCatalogBot(
            token="test-token",
            polling_timeout_seconds=30,
            telegram_client=telegram_client,
            catalog_client=catalog,
        )
        await bot.handle_update(
            TelegramUpdate(
                update_id=1,
                message=TelegramMessage(
                    chat=TelegramChat(id=99),
                    text="Xem https://vn.shp.ee/cMvxmJNm",
                ),
            )
        )

    catalog.import_product.assert_awaited_once_with("https://vn.shp.ee/cMvxmJNm")
    assert sent_payloads[0]["chat_id"] == 99
    assert "Đã thêm sản phẩm" in str(sent_payloads[0]["text"])
    assert "https://s.shopee.vn/an_redir" in str(sent_payloads[0]["text"])


async def test_telegram_message_without_shopee_url_gets_instruction() -> None:
    sent_payloads: list[dict[str, object]] = []

    def telegram_handler(request: httpx.Request) -> httpx.Response:
        sent_payloads.append(json.loads(request.content))
        return httpx.Response(200, json={"ok": True})

    catalog = AsyncMock(spec=CatalogApiClient)
    async with httpx.AsyncClient(
        transport=httpx.MockTransport(telegram_handler)
    ) as telegram_client:
        bot = TelegramCatalogBot(
            token="test-token",
            polling_timeout_seconds=30,
            telegram_client=telegram_client,
            catalog_client=catalog,
        )
        await bot.handle_update(
            TelegramUpdate(
                update_id=1,
                message=TelegramMessage(chat=TelegramChat(id=99), text="xin chào"),
            )
        )

    catalog.import_product.assert_not_awaited()
    assert "link sản phẩm Shopee" in str(sent_payloads[0]["text"])


async def test_telegram_excel_imports_unique_links_with_delay(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "https://s.shopee.vn/first"
    worksheet["A2"] = "https://s.shopee.vn/first"
    worksheet["A3"] = "Xem https://vn.shp.ee/second"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    excel_content = buffer.getvalue()
    sent_messages: list[str] = []

    def telegram_handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/getFile"):
            return httpx.Response(
                200,
                json={
                    "ok": True,
                    "result": {"file_path": "documents/products.xlsx"},
                },
            )
        if request.url.path.endswith("/documents/products.xlsx"):
            return httpx.Response(200, content=excel_content)
        payload = json.loads(request.content)
        sent_messages.append(str(payload["text"]))
        return httpx.Response(200, json={"ok": True})

    product = AffiliateProductResponse(
        shop_id="123",
        item_id="456",
        title="Sản phẩm",
        image_url="https://down-vn.img.susercontent.com/file/image",
        product_url="https://shopee.vn/product/123/456",
        affiliate_url="https://s.shopee.vn/an_redir?affiliate_id=1",
        metadata_source="shopee_open_graph",
    )
    catalog = AsyncMock(spec=CatalogApiClient)
    catalog.import_product.return_value = product
    sleep = AsyncMock()
    monkeypatch.setattr("app.services.telegram_catalog_bot.asyncio.sleep", sleep)

    async with httpx.AsyncClient(
        transport=httpx.MockTransport(telegram_handler)
    ) as telegram_client:
        bot = TelegramCatalogBot(
            token="test-token",
            polling_timeout_seconds=30,
            telegram_client=telegram_client,
            catalog_client=catalog,
            excel_import_delay_seconds=3,
        )
        await bot.handle_update(
            TelegramUpdate(
                update_id=2,
                message=TelegramMessage(
                    chat=TelegramChat(id=99),
                    document=TelegramDocument(
                        file_id="excel-file-id",
                        file_name="products.xlsx",
                        file_size=len(excel_content),
                    ),
                ),
            )
        )

    assert [call.args[0] for call in catalog.import_product.await_args_list] == [
        "https://s.shopee.vn/first",
        "https://vn.shp.ee/second",
    ]
    sleep.assert_awaited_once_with(3)
    assert "Tìm thấy 2 link" in "\n".join(sent_messages)
    assert "Thành công: 2" in sent_messages[-1]
    assert "Thất bại: 0" in sent_messages[-1]


async def test_telegram_rejects_non_xlsx_document() -> None:
    sent_messages: list[str] = []

    def telegram_handler(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        sent_messages.append(str(payload["text"]))
        return httpx.Response(200, json={"ok": True})

    catalog = AsyncMock(spec=CatalogApiClient)
    async with httpx.AsyncClient(
        transport=httpx.MockTransport(telegram_handler)
    ) as telegram_client:
        bot = TelegramCatalogBot(
            token="test-token",
            polling_timeout_seconds=30,
            telegram_client=telegram_client,
            catalog_client=catalog,
        )
        await bot.handle_update(
            TelegramUpdate(
                update_id=3,
                message=TelegramMessage(
                    chat=TelegramChat(id=99),
                    document=TelegramDocument(
                        file_id="csv-file-id",
                        file_name="products.csv",
                    ),
                ),
            )
        )

    catalog.import_product.assert_not_awaited()
    assert sent_messages == ["❌ Chỉ hỗ trợ file Excel định dạng .xlsx."]
